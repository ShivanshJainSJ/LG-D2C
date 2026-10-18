import io
import csv
import pytest
from decimal import Decimal
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.models.user import User, UserRole
from app.models.product import Product, PriceHistory
from app.models.product_import import ProductImportBatch
from app.services.product_import_service import (
    ProductImportService,
    CANONICAL_HEADERS,
    ALLOWED_CATEGORIES_DISPLAY,
)

def create_csv_bytes(headers, rows):
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    for r in rows:
        writer.writerow(r)
    return out.getvalue().encode("utf-8")

def create_xlsx_bytes(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def test_template_generation():
    csv_str = ProductImportService.generate_csv_template()
    assert "sku,product_name,category,description,unit_price,image_url,is_active" in csv_str
    assert "SAMPLE-LG-INV-1000" in csv_str

    xlsx_bytes = ProductImportService.generate_xlsx_template()
    assert len(xlsx_bytes) > 1000
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert headers == CANONICAL_HEADERS

def test_parse_and_validate_new_products(db_session):
    rows = [
        ["00123", "Livguard Pure Inverter 1000", "Inverter", "Sine wave, 12V", "9999.00", "", "true"],
        ["LG-BAT-LITH-1", "Livguard Lithium 100Ah", "Lithium Battery", "Grade A cells, 5-yr warranty", "32500.50", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)
    preview = ProductImportService.validate_and_preview(csv_bytes, "test_new.csv", db_session)

    assert preview["total_rows"] == 2
    assert preview["valid_rows"] == 2
    assert preview["invalid_rows"] == 0
    assert preview["can_import"] is True
    assert preview["rows"][0]["sku"] == "00123"  # Preserved leading zero!
    assert preview["rows"][0]["is_existing_in_db"] is False

def test_in_file_duplicate_sku_rejection(db_session):
    rows = [
        ["LG-DUP-1", "Product 1", "Inverter", "Desc", "5000", "", "true"],
        ["LG-DUP-2", "Product 2", "Inverter", "Desc", "6000", "", "true"],
        ["LG-DUP-1", "Product 1 Duplicate", "Inverter", "Desc", "5000", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)
    preview = ProductImportService.validate_and_preview(csv_bytes, "dups.csv", db_session)

    assert preview["can_import"] is False
    assert preview["file_duplicate_skus_count"] == 1
    # Both rows 2 and 4 should be marked with duplicate errors
    dup_row_1 = [r for r in preview["rows"] if r["row_idx"] == 2][0]
    dup_row_2 = [r for r in preview["rows"] if r["row_idx"] == 4][0]
    assert dup_row_1["is_duplicate_in_file"] is True
    assert dup_row_2["is_duplicate_in_file"] is True
    assert any("Duplicate SKU" in e for e in dup_row_1["errors"])

def test_wrong_header_sequence_rejection(db_session):
    reordered_headers = ["product_name", "sku", "category", "description", "unit_price", "image_url", "is_active"]
    rows = [["Product", "LG-1", "Inverter", "Desc", "1000", "", "true"]]
    csv_bytes = create_csv_bytes(reordered_headers, rows)
    with pytest.raises(ValueError) as excinfo:
        ProductImportService.validate_and_preview(csv_bytes, "wrong_headers.csv", db_session)
    assert "Column header mismatch at position 1" in str(excinfo.value)

def test_invalid_category_and_negative_price(db_session):
    rows = [
        ["LG-INV-INVALID", "Name", "Invalid Category Name", "Desc", "-500", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)
    preview = ProductImportService.validate_and_preview(csv_bytes, "invalid.csv", db_session)

    assert preview["can_import"] is False
    errors = preview["rows"][0]["errors"]
    assert any("Invalid category" in e for e in errors)
    assert any("non-negative" in e for e in errors)

def test_formula_rejection_in_xlsx(db_session):
    rows = [
        ["LG-FORMULA-1", "=SUM(A1:B1)", "Inverter", "Desc", "1000", "", "true"],
    ]
    xlsx_bytes = create_xlsx_bytes(CANONICAL_HEADERS, rows)
    with pytest.raises(ValueError) as excinfo:
        ProductImportService.validate_and_preview(xlsx_bytes, "formulas.xlsx", db_session)
    assert "Formula detected" in str(excinfo.value)

def test_sample_row_rejection(db_session):
    rows = [
        ["SAMPLE-LG-INV-1000", "Livguard Inverter 1000VA (Sample Data - Do Not Import)", "Inverter", "Desc", "12500.00", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)
    preview = ProductImportService.validate_and_preview(csv_bytes, "sample.csv", db_session)
    assert preview["can_import"] is False
    assert any("Sample template row detected" in e for e in preview["rows"][0]["errors"])

def test_skip_existing_policy_execution(db_session):
    admin = db_session.query(User).filter(User.role == UserRole.ADMIN).first()

    # Pre-populate a product
    existing = Product(
        sku="LG-EXISTING-1",
        name="Original Name",
        category="inverter",
        description="Original Desc",
        base_price=Decimal("10000.00"),
        is_active=True
    )
    db_session.add(existing)
    db_session.commit()

    rows = [
        ["LG-EXISTING-1", "Modified Name Attempt", "Inverter", "New Desc", "15000.00", "", "true"],
        ["LG-NEW-IMPORT-1", "New Imported Product", "Lithium Battery", "Desc", "40000.00", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)

    batch = ProductImportService.execute_import(
        file_bytes=csv_bytes,
        filename="skip_test.csv",
        conflict_policy="skip",
        overwrite_confirmed=False,
        current_admin=admin,
        db=db_session
    )

    assert batch.status == "COMPLETED"
    assert batch.inserted_count == 1
    assert batch.skipped_count == 1
    assert batch.updated_count == 0

    # Verify existing product remained unchanged
    db_session.refresh(existing)
    assert existing.name == "Original Name"
    assert existing.base_price == Decimal("10000.00")

    # Verify new product was inserted with price history
    new_prod = db_session.query(Product).filter(Product.sku == "LG-NEW-IMPORT-1").first()
    assert new_prod is not None
    assert new_prod.base_price == Decimal("40000.00")
    hist = db_session.query(PriceHistory).filter(PriceHistory.product_id == new_prod.id).first()
    assert hist is not None
    assert hist.new_price == Decimal("40000.00")

def test_overwrite_unconfirmed_rejected(db_session):
    admin = db_session.query(User).filter(User.role == UserRole.ADMIN).first()
    rows = [
        ["LG-OVERWRITE-1", "Name", "Inverter", "Desc", "10000.00", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)

    with pytest.raises(ValueError) as excinfo:
        ProductImportService.execute_import(
            file_bytes=csv_bytes,
            filename="overwrite_unconfirmed.csv",
            conflict_policy="overwrite",
            overwrite_confirmed=False,
            current_admin=admin,
            db=db_session
        )
    assert "Overwrite confirmation checkbox must be explicitly checked" in str(excinfo.value)

def test_overwrite_confirmed_execution_and_price_history(db_session):
    admin = db_session.query(User).filter(User.role == UserRole.ADMIN).first()

    # Pre-populate product
    existing = Product(
        sku="LG-EXISTING-2",
        name="Old Inverter",
        category="inverter",
        description="Old Inverter Desc",
        base_price=Decimal("12000.00"),
        is_active=True
    )
    db_session.add(existing)
    db_session.commit()

    rows = [
        ["LG-EXISTING-2", "Updated Inverter Pro", "Inverter", "Updated Inverter Pro Desc", "14500.00", "", "true"],
    ]
    csv_bytes = create_csv_bytes(CANONICAL_HEADERS, rows)

    batch = ProductImportService.execute_import(
        file_bytes=csv_bytes,
        filename="overwrite_test.csv",
        conflict_policy="overwrite",
        overwrite_confirmed=True,
        current_admin=admin,
        db=db_session
    )

    assert batch.status == "COMPLETED"
    assert batch.updated_count == 1
    assert batch.inserted_count == 0

    db_session.refresh(existing)
    assert existing.name == "Updated Inverter Pro"
    assert existing.base_price == Decimal("14500.00")

    # Verify price history was recorded for the update
    histories = db_session.query(PriceHistory).filter(PriceHistory.product_id == existing.id).all()
    assert len(histories) > 0
    latest = histories[-1]
    assert latest.old_price == Decimal("12000.00")
    assert latest.new_price == Decimal("14500.00")
    assert "Overwrite" in latest.reason

def test_agent_forbidden_from_import_endpoints(client):
    login_res = client.post("/api/v1/auth/login", json={
        "email": "test_agent@livguard.com",
        "password": "AgentPass123"
    })
    agent_token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {agent_token}"}

    # Template download
    res = client.get("/api/v1/admin/products/import/template/csv", headers=headers)
    assert res.status_code == 403

    # Preview
    res = client.post(
        "/api/v1/admin/products/import/preview",
        headers=headers,
        files={"file": ("test.csv", b"dummy content", "text/csv")}
    )
    assert res.status_code == 403

    # Execute
    res = client.post(
        "/api/v1/admin/products/import/execute",
        headers=headers,
        files={"file": ("test.csv", b"dummy content", "text/csv")},
        data={"conflict_policy": "skip", "overwrite_confirmed": "false"}
    )
    assert res.status_code == 403
